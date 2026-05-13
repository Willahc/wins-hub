#!/usr/bin/env python3
"""Captador ANP — Previsao de Atividades e Investimentos Exploratorios.

NOTA SOBRE ESCOPO (sprint dia 2):
  ANP publica dados agregados em XLSX/PDF anuais ou trimestrais, cada um com
  schema diferente. Para insercao automatica em obras, precisamos mapear
  manualmente cada schema (campos: empresa, bloco, capex_previsto, prazo).

  Esta primeira versao FUNCIONA COMO SCAFFOLD:
    - Tenta CKAN dados.gov.br (retornou 401 no dia 1, retesta).
    - Playwright navega ate a pagina de Previsao de Atividades e Investimentos
      Exploratorios + subpasta de arquivos.
    - Lista arquivos XLSX/CSV.
    - Baixa o(s) mais relevantes, inspeciona sheets/linhas via openpyxl.
    - Loga descobertas, REPORTA STATS mas NAO INSERE obras automaticamente
      (evita lixo no banco com schema variavel).

  Quando schema for validado em sessao dedicada, habilitar INSERT.

STATS_JSON na ultima linha (orchestrator).
"""
from __future__ import annotations

import atexit
import io
import json as _json
import logging
import os
import sys
from typing import Any, Dict, List

sys.path.insert(0, "/app")
sys.path.insert(0, "/app/scripts")

import psycopg2
import requests


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("captar_anp")

FONTE = "anp"
PORTAL_URL = "https://www.gov.br/anp/pt-br/centrais-de-conteudo/dados-abertos/previsao-de-exploracao"
SUBPASTA_ARQUIVOS = "https://www.gov.br/anp/pt-br/centrais-de-conteudo/dados-abertos/arquivos/arquivos-previsao-de-atividades-e-investimentos"

DB_CONFIG = {
    "host": os.getenv("DB_HOST", "db"),
    "port": int(os.getenv("DB_PORT", "5432")),
    "dbname": os.getenv("DB_NAME", "wins_hub"),
    "user": os.getenv("DB_USER", "postgres"),
    "password": os.getenv("DB_PASSWORD", ""),
}

_STATS = {
    "buscados": 0,
    "novos": 0,
    "erros": 0,
    "ckan_ok": 0,
    "xlsx_baixados": 0,
    "sheets_descobertos": 0,
    "linhas_totais": 0,
}


@atexit.register
def _emit_stats():
    print(f"STATS_JSON: {_json.dumps(_STATS)}", flush=True)


def tentar_ckan() -> bool:
    try:
        r = requests.get(
            "https://dados.gov.br/api/3/action/package_search?q=anp&rows=1",
            timeout=15,
            headers={"User-Agent": "Mozilla/5.0"},
        )
    except Exception as e:
        log.warning(f"CKAN HTTP erro: {e}")
        return False
    log.info(f"CKAN status={r.status_code}")
    if r.status_code == 200:
        try:
            return bool(r.json().get("success"))
        except Exception:
            return False
    return False


def coletar_links_xlsx_via_playwright(limite: int = 5) -> List[Dict[str, str]]:
    from playwright.sync_api import sync_playwright

    urls_tentativas = [PORTAL_URL, SUBPASTA_ARQUIVOS]
    coletados: List[Dict[str, str]] = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(
            user_agent=("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                        "(KHTML, like Gecko) Chrome/120 Safari/537.36"),
        )
        page = ctx.new_page()
        for url in urls_tentativas:
            try:
                page.goto(url, timeout=30000, wait_until="domcontentloaded")
            except Exception as e:
                log.warning(f"playwright goto {url}: {e}")
                continue
            items = page.eval_on_selector_all(
                "a",
                "els => els.map(e => ({"
                "  h: e.href, "
                "  t: (e.innerText || '').trim().slice(0, 80)"
                "})).filter(o => /\\.(xlsx|csv|xls)$/i.test(o.h))",
            )
            for it in items:
                if not it.get("h"):
                    continue
                coletados.append({"titulo": (it.get("t") or "")[:200], "url": it["h"]})
            if coletados:
                break
        browser.close()
    # dedup por url
    seen = set()
    uniq: List[Dict[str, str]] = []
    for it in coletados:
        if it["url"] in seen:
            continue
        seen.add(it["url"])
        uniq.append(it)
    return uniq[:limite]


def baixar_xlsx_info(url: str) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        log.warning("openpyxl nao disponivel")
        return {}
    try:
        r = requests.get(url, timeout=120, headers={"User-Agent": "Mozilla/5.0"})
        r.raise_for_status()
    except Exception as e:
        log.warning(f"download {url}: {e}")
        return {}
    try:
        wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    except Exception as e:
        log.warning(f"openpyxl parse {url}: {e}")
        return {}
    info: Dict[str, Any] = {"sheets": [], "total_rows": 0}
    for s in wb.sheetnames:
        ws = wb[s]
        rows = ws.max_row or 0
        cols = ws.max_column or 0
        info["sheets"].append({"name": s, "rows": rows, "cols": cols})
        info["total_rows"] += rows
    return info


def main():
    log.info("ANP — sprint dia 2 (scaffold: CKAN retry + Playwright XLSX descoberta)")

    if tentar_ckan():
        _STATS["ckan_ok"] = 1
        log.info("CKAN respondeu 200 — possivelmente o bloqueio foi transitorio")
    else:
        log.info("CKAN bloqueado (esperado) — fallback Playwright")

    links = coletar_links_xlsx_via_playwright(limite=5)
    _STATS["buscados"] = len(links)
    log.info(f"XLSX candidatos: {len(links)}")
    for li in links:
        log.info(f"  - {li['titulo']!r} -> {li['url'][:120]}")

    conn = psycopg2.connect(**DB_CONFIG)
    try:
        for li in links[:2]:
            info = baixar_xlsx_info(li["url"])
            if not info:
                continue
            _STATS["xlsx_baixados"] += 1
            _STATS["sheets_descobertos"] += len(info["sheets"])
            _STATS["linhas_totais"] += info["total_rows"]
            log.info(f"  + {li['titulo']!r}: {len(info['sheets'])} sheets, "
                     f"{info['total_rows']} linhas totais")
            for sh in info["sheets"][:5]:
                log.info(f"      sheet={sh['name']!r} rows={sh['rows']} cols={sh['cols']}")
    finally:
        conn.close()

    if _STATS["xlsx_baixados"] == 0:
        log.warning("Nenhum XLSX baixado — layout pode ter mudado ou anti-bot")
    else:
        log.info("ANP scaffold OK. INSERT obras requer schema-mapping manual "
                 "(fora deste sprint dia 2 — agendar sessao dedicada).")
    log.info(f"ANP done — {_STATS}")
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
