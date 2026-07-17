#!/usr/bin/env python3
"""Exportacao sob demanda da Planilha Mestre a partir de wins_v2.

Nao reconstrucao a partir da V1. Nao roda automaticamente a cada captura.

Uso:
  python exportar_planilha_mestre_v2.py --out /tmp/planilha.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

# path bootstrap
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if "/app" not in sys.path:
    sys.path.insert(0, "/app")

from services.master_pipeline_v2 import CAMPOS_288, _connect, stable_json  # noqa: E402


def fetch_rows(limit: int | None = None) -> List[Dict[str, Any]]:
    conn = _connect()
    try:
        with conn.cursor() as cur:
            sql = """
                SELECT id, fonte_id, captador_id, id_externo, hash_conteudo,
                       capturado_em, origem_marcador, versao, v1_obra_id,
                       campos_canonicos, payload, namespace
                  FROM wins_v2.capturas_brutas
                 WHERE origem_marcador IS DISTINCT FROM 'HISTORICO_IMPORTADO'
                    OR origem_marcador = 'HISTORICO_IMPORTADO'
                 ORDER BY capturado_em DESC NULLS LAST
            """
            if limit:
                sql += f" LIMIT {int(limit)}"
            cur.execute(sql)
            cols = [d[0] for d in cur.description]
            rows = [dict(zip(cols, r)) for r in cur.fetchall()]

            # nomes de fonte/captador
            cur.execute("SELECT id, nome FROM wins_v2.fontes")
            fontes = dict(cur.fetchall())
            cur.execute("SELECT id, nome FROM wins_v2.captadores")
            captadores = dict(cur.fetchall())
            for r in rows:
                r["_fonte_nome"] = fontes.get(r["fonte_id"])
                r["_captador_nome"] = captadores.get(r["captador_id"])
            return rows
    finally:
        conn.close()


def row_to_288(row: Dict[str, Any]) -> List[Any]:
    campos = row.get("campos_canonicos") or {}
    if isinstance(campos, str):
        campos = json.loads(campos)
    out = []
    for key in CAMPOS_288:
        val = campos.get(key)
        if key == "fonte" and not val:
            val = row.get("_fonte_nome")
        if key == "captador" and not val:
            val = row.get("_captador_nome")
        if key == "id_externo" and not val:
            val = row.get("id_externo")
        if key == "captura_id" and not val:
            val = str(row["v1_obra_id"]) if row.get("v1_obra_id") else None
        if key == "hash_payload" and not val:
            val = row.get("hash_conteudo")
        if key == "payload_original" and val is None:
            val = row.get("payload")
        if isinstance(val, (dict, list)):
            val = stable_json(val)
        out.append(val)
    return out


def exportar(out_path: Path, limit: int | None = None) -> Dict[str, Any]:
    from openpyxl import Workbook

    rows = fetch_rows(limit=limit)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("01_DADOS_MESTRE")
    ws.append(list(CAMPOS_288))
    for row in rows:
        ws.append(row_to_288(row))
    # remove default sheet if present
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {
        "arquivo": str(out_path),
        "registros": len(rows),
        "campos": len(CAMPOS_288),
        "gerado_em": datetime.now(timezone.utc).isoformat(),
    }


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        type=Path,
        default=Path(
            os.environ.get(
                "MASTER_EXPORT_PATH",
                f"/tmp/PLANILHA_MESTRE_V2_EXPORT_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            )
        ),
    )
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args(argv)
    info = exportar(args.out, limit=args.limit)
    print(json.dumps(info, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
